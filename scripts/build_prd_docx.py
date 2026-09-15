from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "101828"
MUTED = "667085"
LIGHT_BORDER = "D7DEE8"
HEADER_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
WHITE = "FFFFFF"


def set_run_font(run, *, size=None, bold=None, color=None, italic=None, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, *, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LIGHT_BORDER)
        borders.append(node)


def table_widths(headers):
    count = len(headers)
    if count == 2:
        if headers[0] in {"指标", "术语", "业务", "能力", "实体组", "测试域", "角色"}:
            return [2100, 7260]
        return [2600, 6760]
    if count == 3:
        return [1700, 3300, 4360]
    if count == 4:
        if headers[0] in {"模块", "一级分组"}:
            return [1900, 1850, 2750, 2860]
        if headers[0] in {"编号", "ID"}:
            return [1300, 2900, 2960, 2200]
        return [1800, 2400, 2580, 2580]
    if count == 5:
        return [1300, 1900, 1900, 1960, 2300]
    return [9360 // count] * count


def set_table_geometry(table, widths):
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_borders(table)


def add_numbering_definition(document, *, bullet):
    numbering = document.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(level_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num = OxmlElement("w:abstractNumId")
    abstract_num.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])


def add_inline_markdown(paragraph, text, default_size=11, default_color=INK, *, direct=True):
    token_re = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    position = 0
    for match in token_re.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            if direct:
                set_run_font(run, size=default_size, color=default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            if direct:
                set_run_font(run, size=default_size, bold=True, color=default_color)
            else:
                run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=max(8.5, default_size - 0.5), color=DARK_BLUE, font="Consolas")
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        if direct:
            set_run_font(run, size=default_size, color=default_color)


def add_callout(document, text):
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before=4, after=10, line=1.2)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CALLOUT_FILL)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    p_pr.append(borders)
    add_inline_markdown(paragraph, text, default_size=10.5)


def add_table(document, rows, *, header=True):
    headers = rows[0]
    table = document.add_table(rows=len(rows), cols=len(headers))
    widths = table_widths(headers)
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.rows[row_index].cells[column_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            set_paragraph_spacing(paragraph, before=0, after=0, line=1.05)
            if column_index == 0 and len(headers) <= 3:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif len(value) < 12 and column_index > 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_markdown(paragraph, value, default_size=9.0 if len(headers) >= 4 else 9.5)
            if header and row_index == 0:
                set_cell_shading(cell, HEADER_FILL)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        if header and row_index == 0:
            set_repeat_table_header(table.rows[row_index])
    set_table_geometry(table, widths)
    after = document.add_paragraph()
    set_paragraph_spacing(after, before=0, after=2, line=1.0)
    return table


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
        "Heading 4": (11, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    code_style = styles.add_style("PRD Code", 1)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code_style.font.size = Pt(9)
    code_style.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.right_indent = Inches(0.18)
    code_style.paragraph_format.space_before = Pt(2)
    code_style.paragraph_format.space_after = Pt(2)
    code_style.paragraph_format.line_spacing = 1.05


def configure_page(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(paragraph, after=0, line=1.0)
    left = paragraph.add_run("鲨鱼妹妹设备运营管理后台")
    set_run_font(left, size=9, bold=True, color=MUTED)
    right = paragraph.add_run("    PRD V1.0")
    set_run_font(right, size=9, color=MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(paragraph, after=0, line=1.0)
    run = paragraph.add_run("内部开发与测试基线  |  第 ")
    set_run_font(run, size=9, color=MUTED)
    add_field(paragraph, "PAGE")
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def add_cover(document):
    spacer = document.add_paragraph()
    set_paragraph_spacing(spacer, after=26, line=1.0)
    kicker = document.add_paragraph()
    set_paragraph_spacing(kicker, after=6, line=1.0)
    run = kicker.add_run("PRODUCT REQUIREMENTS DOCUMENT")
    set_run_font(run, size=10, bold=True, color=BLUE)

    title = document.add_paragraph()
    set_paragraph_spacing(title, after=6, line=1.0)
    run = title.add_run("鲨鱼妹妹设备运营管理后台")
    set_run_font(run, size=25, bold=True, color=INK)

    subtitle = document.add_paragraph()
    set_paragraph_spacing(subtitle, after=20, line=1.0)
    run = subtitle.add_run("产品需求文档与开发 / 测试验收基线")
    set_run_font(run, size=14, color=MUTED)

    metadata = [
        ["文档版本", "V1.0"],
        ["系统版本", "后台演示版 V3.2 / 本地数据库 v8"],
        ["基线日期", "2026-08-14"],
        ["文档状态", "开发与测试基线"],
        ["适用对象", "产品、设计、前端、后端、测试、交付与客户演示人员"],
    ]
    for label, value in metadata:
        paragraph = document.add_paragraph()
        set_paragraph_spacing(paragraph, after=3, line=1.1)
        label_run = paragraph.add_run(f"{label}：")
        set_run_font(label_run, size=10.5, bold=True, color=DARK_BLUE)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10.5, color=INK)
    spacer = document.add_paragraph()
    set_paragraph_spacing(spacer, after=5, line=1.0)
    add_callout(document, "当前系统为纯前端 Demo。本地可实现的业务必须真实写入关联数据；第三方与基础设施能力只做确定性静态模拟。")

    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before=14, after=4, line=1.2)
    run = paragraph.add_run("权威基线")
    set_run_font(run, size=11, bold=True, color=DARK_BLUE)
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, after=6, line=1.2)
    add_inline_markdown(paragraph, "当前可运行的 admin-web 代码、后台 V3.2 需求文档，以及本 PRD 中明确列出的验收标准。", default_size=10.5)
    document.add_page_break()


def parse_table(lines, index):
    rows = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", value) for value in values):
            rows.append(values)
        index += 1
    return rows, index


def build_document(markdown_path, output_path):
    source = Path(markdown_path).read_text(encoding="utf-8")
    lines = source.splitlines()
    document = Document()
    configure_page(document)
    configure_styles(document)
    bullet_num = add_numbering_definition(document, bullet=True)
    decimal_num = add_numbering_definition(document, bullet=False)
    add_cover(document)

    start = next(index for index, line in enumerate(lines) if line.startswith("## 1. Summary"))
    index = start
    in_code = False
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            paragraph = document.add_paragraph(style="PRD Code")
            run = paragraph.add_run(line)
            set_run_font(run, size=9, color=DARK_BLUE, font="Consolas")
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            if rows:
                add_table(document, rows)
            continue
        if stripped.startswith("> "):
            add_callout(document, stripped[2:])
            index += 1
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)) - 1, 3)
            paragraph = document.add_paragraph(style=f"Heading {level}")
            add_inline_markdown(paragraph, heading.group(2), direct=False)
            index += 1
            continue
        if stripped.startswith("- "):
            paragraph = document.add_paragraph()
            apply_numbering(paragraph, bullet_num)
            set_paragraph_spacing(paragraph, after=4, line=1.25)
            add_inline_markdown(paragraph, stripped[2:], direct=False)
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            paragraph = document.add_paragraph()
            apply_numbering(paragraph, decimal_num)
            set_paragraph_spacing(paragraph, after=4, line=1.25)
            add_inline_markdown(paragraph, numbered.group(1), direct=False)
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if not next_line or next_line.startswith(("#", "|", "- ", "> ", "```")) or re.match(r"^\d+\.\s+", next_line):
                break
            paragraph_lines.append(next_line)
            index += 1
        paragraph = document.add_paragraph()
        set_paragraph_spacing(paragraph, after=6, line=1.25)
        add_inline_markdown(paragraph, " ".join(paragraph_lines), direct=False)

    core = document.core_properties
    core.title = "鲨鱼妹妹设备运营管理后台 PRD"
    core.subject = "产品需求文档与开发测试验收基线"
    core.author = "鲨鱼妹妹项目组"
    core.keywords = "PRD, 后台管理, 设备运营, 测试验收"
    core.comments = "基于 admin-web 当前运行系统和后台 V3.2 需求整理"
    document.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_prd_docx.py <input.md> <output.docx>")
    build_document(sys.argv[1], sys.argv[2])
